"""Import des tarifs par type de patient saisis par une structure dans son
fichier de collecte (CHP / CHU / CHR "- avec tarifs assurances.xlsx") : les
onglets ACTES et MEDOC y portent, à la suite des colonnes du catalogue, les
colonnes "Tarif non assuré (FCFA)" puis des blocs "Assurance privée n°X :
<nom>" / Tarif privé / PBR habituel / PBR alternatif (patron, 2026-10-05 :
tout sur la même ligne, rempli en même temps que le prix de votre centre).

  - tarif non assuré par acte / produit     -> prix_non_assure_actes
  - tarif privé + PBR par acte x compagnie  -> pbr_complementaires

Les colonnes sont repérées par leur EN-TÊTE (ligne 1), l'onglet des
médicaments par son nom (contient "medoc"/"medic"/"produit"). Rapprochement
par le NOM exact de la colonne B (celui du catalogue, code compris). Par
défaut : simulation (rien n'est écrit), avec le détail des lignes posées et
ignorées. --appliquer pour écrire.

Usage :
    python importer_tarifs_assurances.py <fichier.xlsx> <structure_id> [--appliquer]
"""
import io
import os
import re
import sys

os.environ['PYTHONIOENCODING'] = 'utf-8'
if (getattr(sys.stdout, 'encoding', '') or '').lower().replace('-', '') != 'utf8':
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from openpyxl import load_workbook  # noqa: E402

PLACEHOLDER = 'écrire son nom ici'


def nombre(v):
    if v is None or v == '':
        return None
    try:
        x = float(str(v).replace(' ', '').replace(',', '.'))
    except ValueError:
        return None
    return x if x > 0 else None


def entete(ws, c):
    v = ws.cell(row=1, column=c).value
    return str(v).strip() if v is not None else ''


def analyser_onglet(ws):
    """-> None si l'onglet n'a pas la colonne "Tarif non assuré", sinon
    (col_tarif_na, blocs: [(col_nom, nom_assurance|None, col_tarif, col_pbr1, col_pbr2)])."""
    col_na = None
    blocs = []
    for c in range(1, ws.max_column + 1):
        h = entete(ws, c).lower()
        if h.startswith('tarif non assur'):
            col_na = c
            continue
        # ⭐ Un bloc assurance se reconnaît à sa STRUCTURE (en-tête suivi de
        # "Tarif privé" / "PBR habituel" / "PBR alternatif"), pas à son texte :
        # la structure peut remplacer tout l'en-tête par le seul nom de
        # l'assurance ("GTA"), ou écrire le nom après les deux-points.
        suivants = [entete(ws, c + j).lower() for j in (1, 2, 3)]
        if (suivants[0].startswith('tarif priv') and suivants[1].startswith('pbr habituel')
                and suivants[2].startswith('pbr alternatif')):
            brut = entete(ws, c)
            nom = brut
            if ':' in brut:
                nom = brut.split(':', 1)[1].strip()
            if not nom or PLACEHOLDER in nom.lower() or nom.lower().startswith('assurance privée') or nom.lower().startswith('assurance privee'):
                nom = None
            blocs.append((c, nom, c + 1, c + 2, c + 3))
    if col_na is None:
        return None
    return col_na, blocs


def lire_onglet(ws, col_na, blocs):
    """-> (lignes: [{nom, prix_non_assure, par_compagnie: {nom: (tarif, pbr1, pbr2)}}], ignorees)"""
    lignes, ignorees = [], []
    for r in range(2, ws.max_row + 1):
        nom = ws.cell(row=r, column=2).value
        if nom is None or str(nom).strip() == '':
            continue
        nom = str(nom).strip()
        prix_na = nombre(ws.cell(row=r, column=col_na).value)
        par_compagnie = {}
        for col_nom, compagnie, ct, cp1, cp2 in blocs:
            tarif, pbr1, pbr2 = (nombre(ws.cell(row=r, column=c).value) for c in (ct, cp1, cp2))
            if tarif is None and pbr1 is None and pbr2 is None:
                continue
            if compagnie is None:
                ignorees.append((r, nom, f"bloc colonne {ws.cell(row=1, column=col_nom).coordinate}", "nom de l'assurance non renseigné dans l'en-tête"))
                continue
            if tarif is None and pbr1 is None:
                ignorees.append((r, nom, compagnie, 'PBR alternatif sans PBR habituel ni tarif privé'))
                continue
            par_compagnie[compagnie] = (tarif, pbr1, pbr2)
        if prix_na is None and not par_compagnie:
            continue
        lignes.append({'nom': nom, 'prix_non_assure': prix_na, 'par_compagnie': par_compagnie})
    return lignes, ignorees


def main():
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    if len(args) < 2:
        print(__doc__)
        sys.exit(1)
    chemin, structure_id = args[0], int(args[1])
    appliquer = '--appliquer' in sys.argv

    wb = load_workbook(chemin, data_only=True)
    resume = {}
    for ws in wb.worksheets:
        analyse = analyser_onglet(ws)
        if analyse is None:
            continue
        col_na, blocs = analyse
        type_article = 'produit' if re.search(r'medoc|medic|produit', ws.title, re.I) else 'acte'
        lignes, ignorees = lire_onglet(ws, col_na, blocs)
        resume.setdefault(type_article, []).append((ws.title, lignes, ignorees))
        n_na = sum(1 for l in lignes if l['prix_non_assure'] is not None)
        n_pc = sum(len(l['par_compagnie']) for l in lignes)
        print(f'\n=== onglet {ws.title!r} ({type_article}) : assurances = {[b[1] for b in blocs if b[1]] or "aucune"}')
        print(f'    {n_na} tarif(s) non assuré, {n_pc} tarif(s)/PBR par assurance, {len(ignorees)} ligne(s) ignorée(s)')
        for r, nom, comp, motif in ignorees[:30]:
            print(f'    - ligne {r} « {nom} » [{comp}] : {motif}')
        for l in lignes[:8]:
            print(f'    ex. « {l["nom"]} » : tarif non assuré={l["prix_non_assure"]} ; {l["par_compagnie"]}')
    if not resume:
        print("Aucun onglet avec une colonne 'Tarif non assuré' trouvé — est-ce bien le fichier avec les colonnes tarifs ?")
        sys.exit(1)

    if not appliquer:
        print('\nSimulation terminée — rien écrit. Relancez avec --appliquer pour enregistrer.')
        return

    from app import app  # noqa: E402  (connexion base)
    from models import db, PbrComplementaire  # noqa: E402
    from services.hospitalisation_service import enregistrer_prix_non_assure  # noqa: E402

    with app.app_context():
        poses_na = poses_pc = 0
        for type_article, onglets in resume.items():
            for _, lignes, _ in onglets:
                for l in lignes:
                    if l['prix_non_assure'] is not None:
                        enregistrer_prix_non_assure(structure_id, l['nom'], l['prix_non_assure'],
                                                    user_nom='import fichier', type_article=type_article)
                        poses_na += 1
                    for compagnie, (tarif, pbr1, pbr2) in l['par_compagnie'].items():
                        existante = PbrComplementaire.query.filter_by(
                            structure_id=structure_id, type=type_article, nom_acte=l['nom'], compagnie=compagnie
                        ).first()
                        if existante:
                            existante.tarif_prive, existante.pbr_1, existante.pbr_2 = tarif, pbr1, pbr2
                        else:
                            db.session.add(PbrComplementaire(
                                structure_id=structure_id, type=type_article, nom_acte=l['nom'], compagnie=compagnie,
                                tarif_prive=tarif, pbr_1=pbr1, pbr_2=pbr2, created_by='import fichier'))
                        poses_pc += 1
        db.session.commit()
        print(f'\nOK structure {structure_id} : {poses_na} tarif(s) non assuré et {poses_pc} tarif(s)/PBR par assurance enregistrés.')


if __name__ == '__main__':
    main()
